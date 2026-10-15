"""
test_docx_html_handler.py — the DOCX and HTML handlers, the output fixes and the image flag
(spec D6, D5, sub-step 5).

Fixtures are built at test time: DOCX files with python-docx (footnotes and
endnotes are added as ``word/footnotes.xml`` / ``word/endnotes.xml`` parts, which
python-docx has no API for), HTML as text. The golden markdown is committed under
``tests/fixtures/golden/`` with the DOCX handler version that produced it.

Run ONLY these tests (the large-DOCX timing test runs only when ``slow`` is
selected, or with ``PARSER_RUN_SLOW=1``):
    uv run pytest tests/test_docx_html_handler.py -m "not slow"
    uv run pytest tests/test_docx_html_handler.py -m slow

Regenerate the golden files after a deliberate output change (and a bump of
``HANDLER_VERSIONS["docx"]``):
    PARSER_UPDATE_GOLDEN=1 uv run pytest tests/test_docx_html_handler.py -k golden
"""

from __future__ import annotations

import base64
import io
import json
import os
import socket
import subprocess
import sys
import time
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from markdown_it import MarkdownIt
from PIL import Image

from parser_service import page_pipeline
from parser_service.formats.guards import check_docx
from parser_service.formats.handlers import HANDLER_VERSIONS, parse_unit
from parser_service.parsing import document as markdown_pipeline
from parser_service.parsing.document import _separate_lists_from_tables, parse_to_markdown

GOLDEN_DIR = Path(__file__).parent / "fixtures" / "golden"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NOTES_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.{}s+xml"


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _png(width: int, height: int) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), "teal").save(buf, "PNG")
    buf.seek(0)
    return buf


def _table(doc: DocxDocument, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text


def _note_reference(paragraph: Any, tag: str, note_id: int) -> None:
    """Append a ``w:footnoteReference`` / ``w:endnoteReference`` run to a paragraph."""
    run = OxmlElement("w:r")
    reference = OxmlElement(f"w:{tag}Reference")
    reference.set(qn("w:id"), str(note_id))
    run.append(reference)
    paragraph._p.append(run)


def _add_notes_part(doc: DocxDocument, tag: str, notes: list[str]) -> None:
    """Add ``word/<tag>s.xml`` with Word's two separator notes and ``notes`` (ids 1..n).

    python-docx writes the relationship and the content-type override on save.
    """
    body = [
        f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>',
        f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        f"<w:continuationSeparator/></w:r></w:p></w:{tag}>",
    ]
    for note_id, text in enumerate(notes, 1):
        body.append(
            f'<w:{tag} w:id="{note_id}"><w:p><w:r><w:t xml:space="preserve">{text}</w:t>'
            f"</w:r></w:p></w:{tag}>"
        )
    blob = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:{tag}s xmlns:w="{_W_NS}">{"".join(body)}</w:{tag}s>'
    ).encode()
    part = Part(
        PackURI(f"/word/{tag}s.xml"), _NOTES_CONTENT_TYPE.format(tag), blob, doc.part.package
    )
    doc.part.relate_to(part, RT.FOOTNOTES if tag == "footnote" else RT.ENDNOTES)


def _build_headings(doc: DocxDocument) -> None:
    doc.add_heading("Quarterly report", 0)
    doc.add_heading("Summary", 1)
    paragraph = doc.add_paragraph("Revenue grew in every region. ")
    paragraph.add_run("Costs fell.").bold = True
    doc.add_heading("Regions", 2)
    doc.add_paragraph("The north led the growth.")
    doc.add_heading("Outlook", 3)
    doc.add_paragraph("Stable.")


def _build_lists(doc: DocxDocument) -> None:
    doc.add_heading("Actions", 1)
    doc.add_paragraph("Hire two engineers", style="List Bullet")
    doc.add_paragraph("Open the Perth office", style="List Bullet")
    doc.add_paragraph("Lease signed", style="List Bullet 2")
    _table(doc, [["Owner", "Due"], ["Ana", "May"]])  # directly after a bullet list
    doc.add_paragraph("Plan the budget", style="List Number")
    doc.add_paragraph("Approve the budget", style="List Number")
    _table(doc, [["Step", "Status"], ["1", "done"]])  # directly after a numbered list
    doc.add_paragraph("A closing paragraph.")


def _build_tables(doc: DocxDocument) -> None:
    doc.add_paragraph("Figures by region:")
    _table(doc, [["Region", "Q1", "Q2"], ["North", "10", "12"], ["South", "7", "9"]])
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Merged heading"
    for c, text in enumerate(["a", "b", "c"]):
        table.cell(1, c).text = text
    doc.add_paragraph("End of the figures.")


def _build_footnotes(doc: DocxDocument) -> None:
    section = doc.sections[0]
    section.header.paragraphs[0].text = "HEADER NOISE"
    section.footer.paragraphs[0].text = "FOOTER NOISE"
    doc.add_heading("Sources", 1)
    _note_reference(doc.add_paragraph("The figure is audited."), "footnote", 1)
    _note_reference(doc.add_paragraph("The method is standard."), "footnote", 2)
    _note_reference(doc.add_paragraph("See the appendix."), "endnote", 1)
    _add_notes_part(doc, "footnote", ["Audited by Example and Co.", "ISO 1234, clause 5."])
    _add_notes_part(doc, "endnote", ["The appendix is published separately."])


GOLDEN_CASES: dict[str, Callable[[DocxDocument], None]] = {
    "headings": _build_headings,
    "lists": _build_lists,
    "tables": _build_tables,
    "footnotes": _build_footnotes,
}


def _docx(tmp_path: Path, build: Callable[[DocxDocument], object], name: str = "doc.docx") -> Path:
    doc = Document()
    build(doc)
    path = tmp_path / name
    doc.save(str(path))
    return path


def _markdown(kind: str, path: Path) -> str:
    state, reason, _code, markdown, extra = parse_unit(kind, path, 1, False)
    assert (state, reason) == ("ok", None), reason
    assert extra["route"] == "docling-kept"
    return markdown


def _golden_path(case: str) -> Path:
    return GOLDEN_DIR / f"docx_{case}.md"


def _read_golden(case: str) -> tuple[str, str]:
    """``(handler version, markdown)`` of a golden file (first line: the version).

    The file ends with one newline the markdown does not have: the repository's
    end-of-file hook adds it to every text file, so it is written and dropped here.
    """
    first, _, body = _golden_path(case).read_text(encoding="utf-8").partition("\n")
    prefix, suffix = "<!-- handler_version: ", " -->"
    assert first.startswith(prefix) and first.endswith(suffix), first
    return first[len(prefix) : -len(suffix)], body.removesuffix("\n")


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("case", list(GOLDEN_CASES))
def test_docx_golden(tmp_path: Path, case: str) -> None:
    """The DOCX handler's output equals the golden file of the current handler version.

    An output change with an unchanged ``HANDLER_VERSIONS["docx"]`` fails here: bump
    the version (so finished DOCX work is re-parsed) and regenerate the golden files.
    """
    markdown = _markdown("docx", _docx(tmp_path, GOLDEN_CASES[case]))
    version = HANDLER_VERSIONS["docx"]
    if os.environ.get("PARSER_UPDATE_GOLDEN") == "1":
        GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
        _golden_path(case).write_text(
            f"<!-- handler_version: {version} -->\n{markdown}\n", encoding="utf-8"
        )
    golden_version, golden = _read_golden(case)
    assert golden_version == version, (
        f"HANDLER_VERSIONS['docx'] is {version} but the golden {case} output is from "
        f"version {golden_version}: regenerate it (PARSER_UPDATE_GOLDEN=1)"
    )
    assert markdown == golden, (
        f"the DOCX {case} output changed under handler version {version}: bump "
        "HANDLER_VERSIONS['docx'] and regenerate the golden files"
    )
    if case == "footnotes":
        assert markdown.endswith(
            "## Footnotes\n\nAudited by Example and Co.\n\nISO 1234, clause 5.\n\n"
            "The appendix is published separately."
        )
        assert "NOISE" not in markdown  # headers and footers stay out


def test_list_followed_by_table_gets_a_blank_line(tmp_path: Path) -> None:
    """Fix (a): a blank line between a list line (or its continuation) and a table line."""
    table = "| a | b |\n|---|---|\n| 1 | 2 |"
    for item in ("- x", "* x", "+ x", "1. x", "12) x", "- x\n  continued", "  - nested"):
        assert _separate_lists_from_tables(f"{item}\n{table}") == f"{item}\n\n{table}"
    # Docling indents a table's first row after a nested item: it moves back too.
    indented_first_row = "- x\n    - nested\n    | a | b |\n|---|---|\n| 1 | 2 |"
    assert _separate_lists_from_tables(indented_first_row) == f"- x\n    - nested\n\n{table}"
    unchanged = [
        f"- x\n\n{table}",  # already separated
        "- x\n  | a |\n  |---|\n  | 1 |",  # a table inside the item (every row indented)
        f"text\n{table}",  # not a list
        f"- x\nlazy text\n{table}",  # a non-indented line ends the list
        f"```\n- x\n{table}\n```",  # inside a code fence
        "-x\n| a |",  # not a list marker
    ]
    for text in unchanged:
        assert _separate_lists_from_tables(text) == text

    # Both the handler and the batch parse_to_markdown carry the fix.
    path = _docx(tmp_path, _build_lists)
    for markdown in (_markdown("docx", path), parse_to_markdown(path, escalate=False)["markdown"]):
        assert "- Lease signed\n\n| Owner" in markdown
        assert "2. Approve the budget\n\n|   Step" in markdown


def test_tables_render_as_tables_not_inside_list_items(tmp_path: Path) -> None:
    """CommonMark (plus tables) renders a table after a list as a table, not list text."""
    md = MarkdownIt("commonmark").enable("table")
    unfixed = "- item\n| A | B |\n|---|---|\n| 1 | 2 |"
    assert "<table>" not in md.render(unfixed)  # the bug the fix removes

    for markdown in (
        _markdown("docx", _docx(tmp_path, _build_lists)),
        _separate_lists_from_tables(unfixed),
    ):
        tokens = md.parse(markdown)
        tables = [t for t in tokens if t.type == "table_open"]
        assert tables, markdown
        # Every table is at the top level, never nested in a list item.
        assert all(t.level == 0 for t in tables)
    html = md.render(_markdown("docx", _docx(tmp_path, _build_lists)))
    assert html.count("<table>") == 2 and "<td>Ana</td>" in html


_ENTITY_CHILD = """
import resource, sys, time
from parser_service.formats.handlers import parse_unit
parse_unit("docx", sys.argv[1], 1, False)  # warm the converter: import and model load
before = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
started = time.perf_counter()
state, reason, code, markdown, extra = parse_unit("docx", sys.argv[2], 1, False)
print(state, code, time.perf_counter() - started,
      (resource.getrusage(resource.RUSAGE_SELF).ru_maxrss - before) / 1024, markdown.count("lol"))
"""


def test_entity_expansion_docx_stays_harmless(tmp_path: Path) -> None:
    """A nested-entity DTD (billion laughs) in document.xml never expands.

    Run in a fresh process, after a warm-up parse, so the RSS growth is the bomb's
    own: under 10 s, under 200 MB, and no expanded text in the output.
    """
    plain = _docx(tmp_path, lambda doc: doc.add_paragraph("PLACEHOLDER"), "plain.docx")
    with zipfile.ZipFile(plain) as source:
        parts = {info.filename: source.read(info.filename) for info in source.infolist()}
    entities = ['<!ENTITY lol "lollollollollollollollollollol">']
    entities.append('<!ENTITY lol1 "' + "&lol;" * 10 + '">')
    entities += [f'<!ENTITY lol{i} "' + f"&lol{i - 1};" * 10 + '">' for i in range(2, 10)]
    head, sep, rest = parts["word/document.xml"].decode().partition("?>")
    parts["word/document.xml"] = (
        (head + sep + f"<!DOCTYPE w:document [{''.join(entities)}]>" + rest)
        .replace("PLACEHOLDER", "before &lol9; after")
        .encode()
    )
    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in parts.items():
            target.writestr(name, data)
    assert check_docx(bomb, page_pipeline.docx_guard_limits()) is None  # the guard lets it by

    done = subprocess.run(
        [sys.executable, "-c", _ENTITY_CHILD, str(plain), str(bomb)],
        capture_output=True,
        text=True,
        timeout=120,
        check=True,
    )
    state, code, seconds, rss_growth_mb, lols = done.stdout.split()[-5:]
    assert state in ("ok", "failed"), done.stdout
    if state == "failed":
        assert code == "docling_failed"  # a record, not a dead worker or a hang
    assert float(seconds) < 10.0
    assert float(rss_growth_mb) < 200.0
    assert int(lols) == 0


def test_html_conversion_makes_no_network_request(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Remote ``img``, ``link``, ``script``, ``iframe`` and ``cid:`` sources are never fetched.

    Every way to open a connection or resolve a name is patched to record and raise.
    """
    attempts: list[tuple[str, tuple[Any, ...]]] = []

    def _refuse(name: str) -> Callable[..., Any]:
        def _raise(*args: Any, **kwargs: Any) -> Any:
            attempts.append((name, args))
            raise OSError(f"network disabled in this test ({name})")

        return _raise

    monkeypatch.setattr(socket.socket, "connect", _refuse("connect"))
    monkeypatch.setattr(socket.socket, "connect_ex", _refuse("connect_ex"))
    monkeypatch.setattr(socket, "create_connection", _refuse("create_connection"))
    monkeypatch.setattr(socket, "getaddrinfo", _refuse("getaddrinfo"))
    monkeypatch.setattr(socket, "gethostbyname", _refuse("gethostbyname"))
    monkeypatch.setattr(markdown_pipeline, "_kind_converters", {})  # a fresh HTML converter

    page = tmp_path / "page.html"
    page.write_text(
        "<!doctype html><html><head>"
        '<link rel="stylesheet" href="http://example.invalid/site.css">'
        '<script src="https://example.invalid/app.js"></script>'
        "</head><body><h1>Notice</h1><p>Body text.</p>"
        '<img src="http://example.invalid/remote.png" alt="remote">'
        '<img src="//example.invalid/protocol-relative.png">'
        '<img src="cid:part1.abc@example.invalid">'
        '<img src="file:///etc/hosts">'
        '<iframe src="https://example.invalid/frame"></iframe>'
        "</body></html>",
        encoding="utf-8",
    )
    state, reason, _code, markdown, extra = parse_unit("html", page, 1, False)
    assert (state, reason, extra["route"]) == ("ok", None, "docling-kept")
    assert "# Notice" in markdown and "Body text." in markdown
    assert extra["flags"] == []  # remote, cid: and local images do not count
    assert attempts == []

    # The converter is HTML-only, with every fetch option off.
    converter = markdown_pipeline._document_converter("html")
    [(input_format, option)] = converter.format_to_options.items()
    assert input_format.value == "html"
    options = option.backend_options
    assert (options.enable_remote_fetch, options.enable_local_fetch, options.fetch_images) == (
        False,
        False,
        False,
    )


def test_images_not_ocrd_flag_keeps_the_unit_ok(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An embedded image of at least PARSER_IMAGES_MIN_PIXELS flags the unit; it stays ok.

    A DOCX body picture and an HTML ``data:`` image count; a smaller image, a
    header (furniture) picture, and remote or ``cid:`` images do not.
    """

    def _big_picture(doc: DocxDocument) -> None:
        doc.add_paragraph("A chart follows.")
        doc.add_picture(_png(300, 200))  # 60,000 pixels

    def _small_pictures(doc: DocxDocument) -> None:
        doc.add_paragraph("A logo follows.")
        doc.add_picture(_png(100, 100))  # 10,000 pixels
        doc.sections[0].header.paragraphs[0].add_run().add_picture(_png(400, 400))

    def _html(width: int, height: int) -> Path:
        data = base64.b64encode(_png(width, height).getvalue()).decode()
        path = tmp_path / f"img_{width}x{height}.html"
        path.write_text(
            f'<html><body><p>Scan below.</p><img src="data:image/png;base64,{data}">'
            '<img src="http://example.invalid/huge.png"><img src="cid:big@x"></body></html>',
            encoding="utf-8",
        )
        return path

    cases = [
        ("docx", _docx(tmp_path, _big_picture, "big.docx"), ["images_not_ocrd"]),
        ("docx", _docx(tmp_path, _small_pictures, "small.docx"), []),
        ("html", _html(300, 200), ["images_not_ocrd"]),
        ("html", _html(100, 100), []),
    ]
    for kind, path, flags in cases:
        state, reason, _code, _markdown, extra = parse_unit(kind, path, 1, False)
        assert (state, reason, extra["route"], extra["flags"]) == (
            "ok",
            None,
            "docling-kept",
            flags,
        ), path.name
    # The threshold is read at call time: at 10,000 pixels the small DOCX flags too.
    monkeypatch.setenv("PARSER_IMAGES_MIN_PIXELS", "10000")
    assert parse_unit("docx", cases[1][1], 1, False)[4]["flags"] == ["images_not_ocrd"]
    monkeypatch.delenv("PARSER_IMAGES_MIN_PIXELS")

    # Through the page pipeline (docx enabled as Group 8 will do it): the record and
    # the file status carry the flag, the unit is ok and the file complete.
    monkeypatch.setenv("PARSER_ENABLED_KINDS", "pdf,docx")
    result = page_pipeline.parse_pdf_pages(
        cases[0][1], tmp_path / "store", workers=0, escalate=False
    )
    assert (result["state"], result["kind"], result["flags"]) == (
        "complete",
        "docx",
        ["images_not_ocrd"],
    )
    status = json.loads(Path(result["status_path"]).read_text())
    [unit] = status["pages"]
    assert (status["unit"], unit["state"], unit["route"]) == ("document", "ok", "docling-kept")


def _slow_selected(config: pytest.Config) -> bool:
    markexpr = str(config.getoption("markexpr") or "")
    return os.environ.get("PARSER_RUN_SLOW") == "1" or (
        "slow" in markexpr and "not slow" not in markexpr
    )


@pytest.mark.slow
def test_large_docx_timing(tmp_path: Path, request: pytest.FixtureRequest) -> None:
    """A 600-page synthetic DOCX passes the guard and parses within PARSER_DEADLINE_DOCX_S.

    About a minute, so it runs only when ``slow`` is selected (``-m slow``) or with
    ``PARSER_RUN_SLOW=1``.
    """
    if not _slow_selected(request.config):
        pytest.skip("slow: run with -m slow or PARSER_RUN_SLOW=1")

    def _build(doc: DocxDocument) -> None:
        sentence = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. " * 3
        for page in range(1, 601):
            doc.add_heading(f"Section {page}", 2)
            for i in range(20):
                doc.add_paragraph(f"Paragraph {i} of page {page}. {sentence}")
            if page % 5 == 0:
                doc.add_paragraph("A list item before a table", style="List Bullet")
                _table(doc, [[f"r{r}c{c}" for c in range(4)] for r in range(4)])
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    path = _docx(tmp_path, _build, "large.docx")
    assert check_docx(path, page_pipeline.docx_guard_limits()) is None
    started = time.perf_counter()
    state, reason, _code, markdown, extra = parse_unit("docx", path, 1, False)
    elapsed = time.perf_counter() - started
    assert (state, reason, extra["route"]) == ("ok", None, "docling-kept")
    assert "Paragraph 19 of page 600." in markdown
    assert elapsed < page_pipeline.deadline_for("docx"), elapsed
