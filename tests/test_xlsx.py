"""
test_xlsx.py — XLSX as a page-pipeline kind: the cost guard (``guards.check_xlsx``:
the ZIP index, then the cells and used-range area of the worksheets), the handler
(one section per visible sheet, the ``hidden_sheets_skipped`` and
``formulas_not_calculated`` flags) and the pipeline end to end.

Workbooks are built at test time with openpyxl in ``tmp_path``. The handler and
pipeline tests run Docling on small workbooks; no AWS call is made.

Run ONLY these tests:
    uv run pytest tests/test_xlsx.py
"""

from __future__ import annotations

import io
import json
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from parser_service.formats import handlers as unit_handlers
from parser_service.formats.guards import (
    XlsxLimits,
    check_xlsx,
    count_sheet_cells,
    xlsx_stats,
)
from parser_service.formats.handlers import parse_unit
from parser_service.formats.result import Reply
from parser_service.parsing.document import workbook_sheets
from parser_service.pipeline import api, config, intake, units
from parser_service.storage.store import LocalStore

FIXTURES = Path(__file__).parent / "fixtures"
LIMITS = XlsxLimits(300, 200, 100, 10_000, 500_000, 5_000_000)
_REJECT_BUDGET_S = 0.5


def _workbook(path: Path, rows: list[list[object]], **cells: object) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    for ref, value in cells.items():
        sheet[ref] = value
    book.save(path)
    return path


# ---------------------------------------------------------------------------
# Counting cells, used ranges and formulas, without Docling
# ---------------------------------------------------------------------------


def test_count_sheet_cells_reads_refs_and_formulas_across_chunks() -> None:
    xml = (
        b'<worksheet><cols><col min="1" max="3"/></cols><sheetData>'
        b'<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><f>A1</f><v>3</v></c>'
        b'<c r="C1"><f>A1+B1</f></c></row>'  # a formula with no value
        b'<row r="7"><x:c r="AB7"><x:f t="shared" si="0"/></x:c>'  # prefixed, shared, no value
        b'<c r="D7"><f>1</f><v></v></c><c/></row>'  # an empty value; a cell with no ref
        b'</sheetData><conditionalFormatting><cfRule type="x"/></conditionalFormatting>'
        b"</worksheet>"
    )
    for chunk_size in (7, 64, 1 << 20):  # any chunk boundary gives the same counts
        counts = count_sheet_cells(io.BytesIO(xml), chunk_size=chunk_size)
        assert (counts.cells, counts.max_row, counts.max_col) == (6, 7, 28), chunk_size
        assert counts.formulas_without_value == 3, chunk_size
        assert counts.grid_cells == 7 * 28
        assert not counts.truncated and not counts.limit_exceeded


def test_xlsx_stats_and_workbook_sheets(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.active.title = "R&D <2026>"
    book.active.append([1, 2, "=A1+B1"])
    hidden = book.create_sheet("Lookup")
    hidden.append(["code"])
    hidden.sheet_state = "hidden"
    path = tmp_path / "book.xlsx"
    book.save(path)

    stats = xlsx_stats(path)
    assert (stats.sheets, stats.cells, stats.grid_cells, stats.formulas_without_value) == (
        2,
        4,
        3 + 1,
        1,
    )
    assert workbook_sheets(path) == [("R&D <2026>", "visible"), ("Lookup", "hidden")]


# ---------------------------------------------------------------------------
# The cost guard
# ---------------------------------------------------------------------------


def test_guard_passes_a_normal_workbook_and_refuses_big_ones(tmp_path: Path) -> None:
    assert check_xlsx(FIXTURES / "sheet.xlsx", LIMITS) is None

    many = _workbook(tmp_path / "many.xlsx", [[r * 10 + c for c in range(10)] for r in range(20)])
    small = XlsxLimits(300, 200, 100, 10_000, 150, 5_000_000)
    assert check_xlsx(many, small) == (
        "too_complex",
        "more than 150 cells",
        {"xlsx_max_cells": 150, "xlsx_max_grid_cells": 5_000_000},
    )

    # Two cells whose used range is the whole sheet: Docling never finishes this one.
    corner = _workbook(tmp_path / "corner.xlsx", [], A1="top", XFD1048576="corner")
    started = time.perf_counter()
    rejection = check_xlsx(corner, LIMITS)
    assert time.perf_counter() - started < _REJECT_BUDGET_S
    assert rejection is not None and rejection[0] == "too_complex"
    assert rejection[1] == "used ranges of 17179869184 cells (rows x columns), over 5000000"


def test_guard_refuses_a_zip_bomb(tmp_path: Path) -> None:
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", b" " * (4 << 20))  # 4 MB of spaces
    rejection = check_xlsx(bomb, LIMITS)
    assert rejection is not None
    assert rejection[0] == "zip_bomb" and rejection[1].startswith("ratio ")
    assert rejection[2] == LIMITS.zip_limits()


def test_guard_rejection_is_final_and_never_reaches_the_handler(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    handled: list[Path] = []

    def stub_handler(path: Path, page_number: int, escalate: bool) -> Reply:
        handled.append(path)
        return "ok", None, None, "# parsed", {"route": "docling-kept"}

    monkeypatch.setitem(unit_handlers.HANDLERS, "xlsx", stub_handler)
    store = LocalStore(tmp_path / "store")
    corner = _workbook(tmp_path / "corner.xlsx", [], A1="top", XFD1048576="corner")
    work_id = intake.work_id_for(corner, "xlsx")
    intake._intake(store, work_id, corner, "xlsx")
    record = units._process_page(store, work_id, 1, False, 3, kind="xlsx")
    assert (record["state"], record["reason_code"], handled) == ("failed", "too_complex", [])
    assert record["limits"] == LIMITS.count_limits()

    limits_now = config.current_limits("xlsx")
    assert {**LIMITS.zip_limits(), **LIMITS.count_limits(), "deadline_s": 300}.items() <= (
        limits_now.items()
    )
    assert config._is_final(record, limits_now)
    monkeypatch.setenv("PARSER_XLSX_MAX_GRID_CELLS", str(20_000_000_000))
    assert not config._is_final(record, config.current_limits("xlsx"))


# ---------------------------------------------------------------------------
# The handler and the pipeline (Docling, no AWS)
# ---------------------------------------------------------------------------


def test_handler_writes_one_section_per_visible_sheet(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    revenue = book.active
    revenue.title = "Revenue"
    revenue.append(["Region", "Q1", "Q2", "Total"])
    revenue.append(["North", 10, 20, "=B2+C2"])
    revenue["A5"] = "Notes: figures in AUD"
    costs = book.create_sheet("Costs")
    costs.append(["Item", "Amount"])
    costs.append(["Rent", 1200])
    hidden = book.create_sheet("Lookup")
    hidden.append(["HIDDEN-LOOKUP-MARKER"])
    hidden.sheet_state = "hidden"
    book.create_sheet("Empty")
    path = tmp_path / "report.xlsx"
    book.save(path)

    state, reason, reason_code, markdown, extra = parse_unit("xlsx", path, 1, False)

    assert (state, reason, reason_code, extra["route"]) == ("ok", None, None, "docling-kept")
    assert markdown.startswith("## Revenue\n\n| Region")
    assert "\n\n## Costs\n\n| Item" in markdown
    assert "Notes: figures in AUD" in markdown and "| Notes" not in markdown  # text, not a table
    assert "HIDDEN-LOOKUP-MARKER" not in markdown and "## Lookup" not in markdown
    assert "## Empty" not in markdown
    assert extra["flags"] == ["hidden_sheets_skipped", "formulas_not_calculated"]


def test_empty_workbook_is_an_empty_document(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(path)
    state, _reason, _code, markdown, extra = parse_unit("xlsx", path, 1, False)
    assert (state, markdown, extra["route"], extra["flags"]) == ("ok", "", "empty_document", [])


def test_pipeline_parses_xlsx_and_xlsm(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("PARSER_ENABLED_KINDS", raising=False)
    store = tmp_path / "store"
    macro = tmp_path / "budget.xlsm"
    macro.write_bytes((FIXTURES / "sheet.xlsx").read_bytes())
    for source in (FIXTURES / "sheet.xlsx", macro):
        result = api.parse_pdf_pages(source, store, workers=0, escalate=False)
        assert (result["state"], result["kind"], result["unit"]) == ("complete", "xlsx", "document")
        output = (Path(result["out_dir"]) / "output.md").read_text()
        assert "## Sheet1\n\n| Name" in output and "| Gamma" in output
        manifest = json.loads((Path(result["work_dir"]) / "manifest.json").read_text())
        assert manifest["config"]["kind"] == "xlsx"
        assert "escalation_engine" not in manifest["config"]  # a VLM change never re-parses
