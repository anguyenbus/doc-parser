"""
test_corpus_inventory.py — ``scripts/corpus_inventory.py`` on a local folder.

Offline: the source is a temporary folder with one file per kind, built at test time
(python-docx, ``zipfile``, stdlib ``email``, magic numbers). The S3 path is only
checked for its region refusal, which happens before any client is made.

Run ONLY these tests:
    uv run pytest tests/test_corpus_inventory.py
"""

from __future__ import annotations

import csv
import importlib.util
import json
import struct
import sys
import time
import zipfile
from email.message import EmailMessage
from pathlib import Path
from typing import Any

import docx
import pytest

ROOT = Path(__file__).resolve().parent.parent


def _load_inventory() -> Any:
    """Import scripts/corpus_inventory.py, registered in sys.modules (its dataclasses
    resolve their annotations through it)."""
    spec = importlib.util.spec_from_file_location(
        "corpus_inventory", ROOT / "scripts" / "corpus_inventory.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules["corpus_inventory"] = module
    spec.loader.exec_module(module)
    return module


inventory = _load_inventory()

MARKER = "ZEBRA-MARKER-7731"


def _ole2_msg() -> bytes:
    header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504)
    entries = b""
    for index, name in enumerate(("Root Entry", "__substg1.0_0037001F")):
        raw = (name + "\x00").encode("utf-16-le")
        entry = raw.ljust(64, b"\x00") + struct.pack("<HB", len(raw), 5 if index == 0 else 2)
        entries += entry.ljust(128, b"\x00")
    return header + entries.ljust(512, b"\x00") + MARKER.encode("utf-16-le")


def _corpus(root: Path) -> None:
    root.mkdir()
    document = docx.Document()
    for _ in range(3):
        document.add_paragraph(MARKER)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = MARKER
    document.save(str(root / "letter.docx"))
    document.save(str(root / "renamed.pdf"))  # a DOCX named .pdf: a mismatch
    with zipfile.ZipFile(root / "sheet.xlsx", "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    (root / "report.pdf").write_bytes(b"%PDF-1.7\n" + MARKER.encode() + b"\n%%EOF\n")
    (root / "notes.txt").write_text(f"{MARKER}\nsecond line\n")
    (root / "sub").mkdir()
    (root / "sub" / "page.html").write_text(f"<html><body><p>{MARKER}</p></body></html>")
    message = EmailMessage()
    message["From"] = "a@example.com"
    message["To"] = "b@example.com"
    message["Subject"] = MARKER
    message["Date"] = "Tue, 14 Oct 2026 09:00:00 +1100"
    message.set_content(MARKER)
    (root / "mail.eml").write_bytes(message.as_bytes())
    (root / "outlook.msg").write_bytes(_ole2_msg())
    (root / "data.csv").write_text(f"a,b\n{MARKER},1\n")
    (root / "photo.png").write_bytes(b"\x89PNG\r\n\x1a\n" + bytes(64))


def test_inventory_on_a_local_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    corpus, out_dir, csv_path = tmp_path / "corpus", tmp_path / "report", tmp_path / "files.csv"
    _corpus(corpus)

    code = inventory.main(
        ["--prefix", str(corpus), "--out-dir", str(out_dir), "--csv", str(csv_path)]
    )

    assert code == 0
    report = json.loads((out_dir / "inventory.json").read_text())
    assert report["files"] == 10
    assert report["by_detected_type"] == {
        "docx": 2,
        "xlsx": 1,
        "pdf": 1,
        "text": 2,
        "html": 1,
        "eml": 1,
        "msg": 1,
        "image": 1,
    }
    assert report["by_extension"] == {
        "docx": 1,
        "pdf": 2,
        "xlsx": 1,
        "txt": 1,
        "html": 1,
        "eml": 1,
        "msg": 1,
        "csv": 1,
        "png": 1,
    }
    assert report["by_kind"] == {
        "docx": 2,
        "pdf": 1,
        "txt": 1,
        "html": 1,
        "eml": 1,
        "msg": 1,
        "none": 3,
    }
    size_docx = report["size_bytes_by_kind"]["docx"]
    letter_size = (corpus / "letter.docx").stat().st_size
    assert size_docx == {"n": 2, "p50": letter_size, "p99": letter_size, "max": letter_size}
    assert set(report["size_bytes_by_kind"]) == {"docx", "pdf", "txt", "html", "eml", "msg", "none"}
    assert report["docx"]["files"] == 2
    assert report["docx"]["paragraphs"]["max"] >= 3 + 4
    assert report["docx"]["table_cells"]["max"] == 4
    assert report["docx"]["entries"]["n"] == 2
    assert report["docx"]["ratio"]["p50"] > 1
    assert report["mismatches"] == {"files": 1, "by_pair": {"pdf->docx": 1}}
    assert report["msg_share"]["files"] == 1
    assert report["msg_share"]["file_share"] == pytest.approx(1 / 10)
    assert report["errors"] == {}
    markdown = (out_dir / "inventory.md").read_text()
    assert "# Corpus inventory" in markdown and "## DOCX" in markdown

    # The per-file CSV: keys and numbers only, never document text.
    raw = csv_path.read_text()
    assert MARKER not in raw
    assert MARKER not in json.dumps(report) and MARKER not in markdown
    with csv_path.open() as handle:
        rows = list(csv.DictReader(handle))
    assert tuple(rows[0]) == inventory.CSV_COLUMNS
    assert {row["key"] for row in rows} >= {"sub/page.html", "renamed.pdf", "outlook.msg"}
    letter = next(row for row in rows if row["key"] == "letter.docx")
    assert int(letter["docx_table_cells"]) == 4
    assert int(letter["docx_paragraphs"]) >= 3 + 4

    # --max-files bounds the run; a CSV inside the repository is refused.
    assert (
        inventory.main(
            [
                "--prefix",
                str(corpus),
                "--out-dir",
                str(out_dir),
                "--csv",
                str(csv_path),
                "--max-files",
                "3",
            ]
        )
        == 0
    )
    assert json.loads((out_dir / "inventory.json").read_text())["files"] == 3
    inside = inventory.REPO_ROOT / "inventory.csv"
    assert (
        inventory.main(["--prefix", str(corpus), "--out-dir", str(out_dir), "--csv", str(inside)])
        == 1
    )
    assert not inside.exists()

    # S3 outside ap-southeast-2 is refused before any client is made.
    monkeypatch.setenv("AWS_REGION", "us-east-1")
    monkeypatch.setattr(inventory.aws, "s3_client", pytest.fail)
    assert (
        inventory.main(
            ["--prefix", "s3://bucket/corpus", "--out-dir", str(out_dir), "--csv", str(csv_path)]
        )
        == 1
    )


def test_time_sample_kills_at_the_limit(tmp_path: Path) -> None:
    sleeper = [sys.executable, "-c", "import time; time.sleep(60)"]
    start = time.monotonic()
    killed = inventory.measure_child(sleeper, time_limit_s=0.5, poll_s=0.05)
    assert time.monotonic() - start < 10
    assert killed["verdict"] == "killed_time"
    assert 0.5 <= killed["seconds"] < 10
    assert killed["peak_rss_bytes"] > 0

    finished = inventory.measure_child([sys.executable, "-c", "pass"], time_limit_s=30)
    assert finished["verdict"] == "finished" and finished["returncode"] == 0
    failed = inventory.measure_child([sys.executable, "-c", "raise SystemExit(3)"], 30)
    assert failed["verdict"] == "failed" and failed["returncode"] == 3

    # The sample takes the largest file per kind plus random ones, and times each in
    # the script's own child mode (TXT: no Docling).
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    for index, size in enumerate((10, 5000, 300)):
        (corpus / f"n{index}.txt").write_text("x" * size)
    source = inventory.LocalSource(corpus)
    rows = [inventory.measure_file(source, key, size, tmp_path) for key, size in source.list()]
    assert [row.key for row in inventory.pick_time_sample(rows, 2, seed=0)][0] == "n1.txt"
    inventory.time_sample(source, rows, 1, 0, tmp_path, time_limit_s=60, rss_ceiling_bytes=None)
    timed = [row for row in rows if row.parse_verdict]
    assert [(row.key, row.parse_verdict) for row in timed] == [("n1.txt", "finished")]
    report = inventory.aggregate(rows, str(corpus))
    assert report["timing"]["txt"]["verdicts"] == {"finished": 1}
