"""
make_demo_inputs.py — one small input per kind for the T5 formats test.

    uv run --extra dev python deploy/scripts/make_demo_inputs.py <out dir> <run tag>

Writes demo-report.html, demo-review.docx, demo-memo.txt and demo-mixed.pdf. Every
file carries the run tag, so a new run has new bytes (the store keys work/ by
content) and the tag shows up in the parsed output. The PDF is
tests/fixtures/sample_mixed_10.pdf (10 pages, digital and scanned) with the tag
appended as an incremental metadata update, so its pages are unchanged.
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def write_html(path: Path, tag: str) -> None:
    path.write_text(
        f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Quarterly records report</title></head>
<body>
<h1>Quarterly records report</h1>
<p>Run {tag}. This page is parsed by the page service as one HTML unit.</p>
<h2>Highlights</h2>
<ul>
  <li>All branches filed their retention schedules.</li>
  <li>Eastfield still has a scanning backlog.</li>
</ul>
<h2>Backlog by branch</h2>
<table>
  <tr><th>Branch</th><th>Pages scanned</th><th>Backlog</th></tr>
  <tr><td>Northgate</td><td>12,480</td><td>0</td></tr>
  <tr><td>Eastfield</td><td>3,115</td><td>4,200</td></tr>
</table>
<p>Remote content is never fetched: <img src="https://example.invalid/pixel.png" alt="tracking pixel"></p>
</body></html>
""",
        encoding="utf-8",
    )


def write_docx(path: Path, tag: str) -> None:
    import docx  # python-docx, in the dev extra

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "CONFIDENTIAL - header text, not in the output"
    d.add_heading("Compliance review", 1)
    d.add_paragraph(f"Run {tag}. This Word document is parsed natively, without a split.")
    d.add_paragraph("The scanning backlog at Eastfield is estimated at 4,200 pages.")
    d.add_heading("Actions", 2)
    for item in ("Update the retention schedule", "Train two records officers"):
        d.add_paragraph(item, style="List Bullet")
    table = d.add_table(rows=3, cols=3)  # directly after the list: the blank-line fix
    for r, row in enumerate(
        [
            ["Branch", "Scanned", "Backlog"],
            ["Northgate", "12,480", "0"],
            ["Eastfield", "3,115", "4,200"],
        ]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    tmp = path.with_suffix(".tmp.docx")
    d.save(tmp)
    # python-docx has no footnote API: add one to the backlog sentence by hand.
    notes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes {W}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Estimate from the branch manager, 18 March 2026.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = (
                    data.decode()
                    .replace(
                        "4,200 pages.</w:t></w:r>",
                        '4,200 pages.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r>',
                    )
                    .encode()
                )
            elif item.filename == "word/_rels/document.xml.rels":
                data = (
                    data.decode()
                    .replace(
                        "</Relationships>",
                        '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                        'relationships/footnotes" Target="footnotes.xml"/></Relationships>',
                    )
                    .encode()
                )
            elif item.filename == "[Content_Types].xml":
                data = (
                    data.decode()
                    .replace(
                        "</Types>",
                        '<Override PartName="/word/footnotes.xml" ContentType="application/'
                        'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
                    )
                    .encode()
                )
            zout.writestr(item, data)
        zout.writestr("word/footnotes.xml", notes)
    tmp.unlink()


def write_txt(path: Path, tag: str) -> None:
    memo = f"""MEMO - Scanning backlog
Run: {tag}
To: Records team <records@example.invalid>
Ref: file_name_v2_final

# of boxes: 3
Total pages: 4,200 * approx

Box 1: delivery dockets
Box 2: staff timesheets
    (indented note: some are water damaged)
---
Please reply by Friday."""
    # Windows line endings; the passthrough keeps every line and character as written.
    path.write_bytes(memo.replace("\n", "\r\n").encode("utf-8"))


def write_pdf(path: Path, tag: str) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter(str(REPO / "tests/fixtures/sample_mixed_10.pdf"), incremental=True)
    writer.add_metadata({"/TestCopy": tag})
    with path.open("wb") as handle:
        writer.write(handle)


def main() -> None:
    out, tag = Path(sys.argv[1]), sys.argv[2]
    out.mkdir(parents=True, exist_ok=True)
    write_html(out / "demo-report.html", tag)
    write_docx(out / "demo-review.docx", tag)
    write_txt(out / "demo-memo.txt", tag)
    write_pdf(out / "demo-mixed.pdf", tag)
    for f in sorted(out.glob("demo-*")):
        print(f"{f.name}\t{f.stat().st_size}")


if __name__ == "__main__":
    main()
